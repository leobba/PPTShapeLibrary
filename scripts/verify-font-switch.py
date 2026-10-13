# -*- coding: utf-8 -*-
"""
「一键换字体」COM 语义验证脚本（可复跑）
========================================
验证与 TextToolService.SetFonts 完全相同的调用逻辑是否达到预期：
  1) 中文 -> Font.NameFarEast，英文/数字 -> Font.NameAscii / NameOther
  2) 覆盖原有显式字体（宋体/黑体）与主题字体引用（+mn-*），即"取消原字体"
  3) 覆盖表格单元格内文字
  4) 覆盖组合形状(grpSp)内部文字
用法：python scripts/verify-font-switch.py
退出码：0=通过，1=失败。需要本机安装 PowerPoint。
"""
import os
import re
import sys
import tempfile
import zipfile

from pptx import Presentation
from pptx.util import Inches, Pt
import win32com.client

# Windows 控制台默认 GBK：脚本里的 ✓/✗ 会让 print 抛 UnicodeEncodeError，
# 结果"验证全过却以 exit 1 结束"。统一改 UTF-8 输出，无法编码的字符用替代符。
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

CN_FONT = "微软雅黑"
EN_FONT = "Arial"


def build_source(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # ① 显式原字体（宋体）：中英混排
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(1))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "中文标题 Hello World 2024"
    r.font.name = "宋体"
    r.font.size = Pt(24)

    # ② 主题字体（不显式设置字体 → 走 +mn-ea/+mn-lt）
    tb2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(6), Inches(1))
    r2 = tb2.text_frame.paragraphs[0].add_run()
    r2.text = "主题字体段落 Theme Font 456"
    r2.font.size = Pt(20)

    # ③ 表格内文字（显式原字体 黑体）
    tbl = slide.shapes.add_table(2, 2, Inches(0.5), Inches(3), Inches(6), Inches(1.5)).table
    tbl.cell(0, 0).text = "表格中文 Chart一"
    tbl.cell(0, 1).text = "Cell English 789"
    for rr in range(2):
        for cc in range(2):
            for para in tbl.cell(rr, cc).text_frame.paragraphs:
                for run in para.runs:
                    run.font.name = "黑体"
    prs.save(path)


def apply_fonts_via_com(src, dst):
    """与 TextToolService.SetFonts 相同的递归逻辑（组合 → 表格 → 文本框）。"""
    app = win32com.client.Dispatch("PowerPoint.Application")
    # PowerPoint 的 COM 服务器是单实例：若用户已经开着 PowerPoint，直接 app.Quit()
    # 会把用户的会话一起关掉（未保存的文稿会丢）。因此只在"我们自己是唯一使用者"时退出。
    preexisting = 0
    try:
        preexisting = app.Presentations.Count
    except Exception:
        preexisting = 0
    stats = {"shapes": 0, "chars": 0}

    def set_on_range(rg):
        if rg is None or rg.Length <= 0:
            return
        f = rg.Font
        f.NameFarEast = CN_FONT
        f.NameAscii = EN_FONT
        f.NameOther = EN_FONT
        stats["shapes"] += 1
        stats["chars"] += rg.Length

    def walk(shp, depth=0):
        if shp is None or depth > 6:
            return
        if shp.Type == 6:                      # msoGroup
            for i in range(1, shp.GroupItems.Count + 1):
                walk(shp.GroupItems(i), depth + 1)
            return
        if shp.HasTable:
            t = shp.Table
            for rr in range(1, t.Rows.Count + 1):
                for cc in range(1, t.Columns.Count + 1):
                    try:
                        set_on_range(t.Cell(rr, cc).Shape.TextFrame.TextRange)
                    except Exception:
                        pass
            return
        if shp.HasTextFrame and shp.TextFrame.HasText:
            set_on_range(shp.TextFrame.TextRange)

    try:
        pres = app.Presentations.Open(src, ReadOnly=True, WithWindow=False)
        s = pres.Slides(1)
        # 造一个组合形状，验证组合内部文字
        b1 = s.Shapes.AddTextbox(1, 600, 400, 200, 40)
        b1.TextFrame.TextRange.Text = "组合内中文 GroupText"
        b2 = s.Shapes.AddTextbox(1, 600, 450, 200, 40)
        b2.TextFrame.TextRange.Text = "Group English 321"
        try:
            s.Shapes.Range([s.Shapes.Count - 1, s.Shapes.Count]).Group()
        except Exception as e:
            print("[warn] 组合形状创建失败，跳过组合场景:", e)

        for i in range(1, s.Shapes.Count + 1):
            walk(s.Shapes(i))
        pres.SaveAs(dst)
        pres.Close()
    finally:
        # 只有在我们开始时 PowerPoint 没有打开任何文稿时才退出进程，避免关掉用户正在编辑的文稿
        if preexisting == 0:
            try:
                app.Quit()
            except Exception:
                pass
        else:
            print("[info] 检测到 PowerPoint 已有打开的文稿，跳过 Quit（不干扰用户会话）")
    return stats


def slide_xml(path):
    with zipfile.ZipFile(path) as z:
        return z.read("ppt/slides/slide1.xml").decode("utf-8")


def main():
    work = tempfile.mkdtemp(prefix="fontswitch_")
    src, dst = os.path.join(work, "src.pptx"), os.path.join(work, "out.pptx")
    print("工作目录:", work)
    build_source(src)
    stats = apply_fonts_via_com(src, dst)
    print(f"COM 应用完成: 形状 {stats['shapes']} 个, 约 {stats['chars']} 字")

    xml = slide_xml(dst)
    fails = []

    left_cn = len(re.findall("宋体", xml)) + len(re.findall("黑体", xml))
    theme_refs = len(re.findall(r'typeface="\+m', xml))
    n_ea = len(re.findall(r'typeface="%s"' % CN_FONT, xml))
    n_latin = len(re.findall(r'typeface="%s"' % EN_FONT, xml))

    print("\n--- 校验 ---")
    print(f"原字体残留(宋体/黑体): {left_cn}   [期望 0]")
    print(f"主题字体引用(+m*): {theme_refs}   [期望 0]")
    print(f'a:ea="{CN_FONT}": {n_ea}   [期望 >0]')
    print(f'a:latin="{EN_FONT}": {n_latin}   [期望 >0]')
    if left_cn:
        fails.append("原字体仍有残留")
    if theme_refs:
        fails.append("主题字体引用未覆盖")
    if not n_ea or not n_latin:
        fails.append("新字体未写入")

    # 中英分派：同一 run 内 latin + ea 并存
    for probe in ("中文标题", "主题字体段落", "表格中文", "组合内中文"):
        i = xml.find(probe)
        if i < 0:
            fails.append(f"未找到文本 {probe}")
            continue
        ps = max(xml.rfind("<a:p>", 0, i), xml.rfind("<a:p ", 0, i))
        seg = xml[ps:xml.find("</a:p>", i)]
        latin = re.search(r'<a:latin typeface="([^"]*)"', seg)
        ea = re.search(r'<a:ea typeface="([^"]*)"', seg)
        got = (latin.group(1) if latin else "-", ea.group(1) if ea else "-")
        ok = got == (EN_FONT, CN_FONT)
        print(f"  「{probe}」 latin={got[0]} ea={got[1]}  {'✓' if ok else '✗'}")
        if not ok:
            fails.append(f"{probe} 的中英字体分派不正确")

    print("\nRESULT:", "PASS" if not fails else "FAIL -> " + "; ".join(fails))
    return 0 if not fails else 1


if __name__ == "__main__":
    sys.exit(main())
